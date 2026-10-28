"""UTC/local boundaries and Actual Log regressions from operator examples."""
import io
from datetime import date, datetime, timedelta
from unittest.mock import Mock, patch

from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase, override_settings
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import load_workbook

from core.irrig_time import (complete_irrigation_days, irrigation_day_for_utc,
                            irrigation_day_utc_bounds, runtime_window_bounds,
                            utc_timestamp_to_local)
from core.management.commands.import_maxicom_mdb_linux import actual_runtime_values
from core.models import (MaxicomController, MaxicomFlowDaily, MaxicomFlowZone,
                         MaxicomRuntime, Patch, Zone)

CASES = [
    (7, '202609020654', '202609030654',
     {'7-3-19': 20, '7-4-24': 821, '7-6-24': 852, '7-8-16': 15}),
    (2, '20260827070100', '20260828070159',
     {'2-2-21': 20, '2-3-15': 32, '2-5-2': 20, '2-5-23': 1500,
      '2-6-11': 168, '2-6-15': 84, '2-6-19': 25, '2-6-20': 20, '2-6-21': 30}),
    (44, '20260818070100', '20260819065759',
     {'44-3-1': 30, '44-3-2': 30, '44-3-4': 10, '44-3-20': 10,
      '44-3-21': 15, '44-3-23': 633, '44-4-24': 817}),
]


@override_settings(TIME_ZONE='Asia/Shanghai')
class MaxicomTimeTests(SimpleTestCase):
    def test_operator_windows_convert_to_utc_with_second_precision(self):
        self.assertEqual(runtime_window_bounds('202609020654', '202609030654'),
                         ('20260901225400', '20260902225459'))
        self.assertEqual(runtime_window_bounds('20260818070100', '20260819065759'),
                         ('20260817230100', '20260818225759'))

    def test_local_calendar_and_irrigation_days_cross_year_boundaries(self):
        self.assertEqual(runtime_window_bounds('20260101', '20260101'),
                         ('20251231160000', '20260101155959'))
        self.assertEqual(irrigation_day_utc_bounds(date(2026, 1, 1), date(2026, 1, 1)),
                         ('20251231140000', '20260101135959'))

    def test_utc_labels_and_local_irrigation_day_agree(self):
        self.assertEqual(utc_timestamp_to_local('20260901225400'), '20260902065400')
        self.assertEqual(irrigation_day_for_utc('20260901135959'), date(2026, 9, 1))
        self.assertEqual(irrigation_day_for_utc('20260901140000'), date(2026, 9, 2))

    def test_invalid_and_reversed_windows_do_not_silently_shift(self):
        for start, end in [('20260230', '20260301'), ('bad', '20260902'),
                           ('20260903', '20260902')]:
            with self.assertRaises(ValueError):
                runtime_window_bounds(start, end)

    def test_daily_meter_range_does_not_include_the_previous_day(self):
        self.assertEqual(complete_irrigation_days('202610232200', '202610242159'),
                         (date(2026, 10, 24), date(2026, 10, 24)))
        self.assertIsNone(complete_irrigation_days('202609020654', '202609030654'))

    def test_actual_runtime_excludes_forecast_and_unrelated_signals(self):
        reading = {'XactStamp': '20260901225400', 'XactIndex': '527',
                   'SignalTable': 'S', 'SignalType': 'M', 'SignalValue': '20'}
        rows = [reading, {**reading, 'SignalValue': '0'},
                {**reading, 'SignalValue': '-2'}, {**reading, 'SignalType': 'P'},
                {**reading, 'SignalTable': 'C'}, {**reading, 'XactStamp': 'bad'},
                {'TimeStamps': '20260901225400', 'StationID': '527', 'RunTime': '1'}]
        self.assertEqual(list(actual_runtime_values(rows)), [('20260901225400', 527, 20)])


@override_settings(TIME_ZONE='Asia/Shanghai', CACHES={
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}})
class MaxicomRuntimeViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser('runtime-review', '', 'test-only')
        cls.ccus, cls.stations = {}, {}
        boundary = [{'lat': 31, 'lng': 121}, {'lat': 31.01, 'lng': 121},
                    {'lat': 31, 'lng': 121.01}]
        for ccu_num, start, end, values in CASES:
            ccu, _ = Patch.objects.get_or_create(code=f'CCU{ccu_num}', defaults={'name': f'Area {ccu_num}'})
            cls.ccus[ccu_num] = ccu
            controllers = {}
            for code, minutes in values.items():
                _, sat, channel = map(int, code.split('-'))
                if sat not in controllers:
                    controllers[sat] = MaxicomController.objects.create(
                        site=ccu, mdb_index=900000 + ccu_num * 100 + sat,
                        name=f'SAT {ccu_num}-{sat}', site_number=ccu_num,
                        link_number=1, link_channel=sat)
                station = Patch.objects.create(
                    code=f'station-time-test-{code}', name=code, parent=ccu,
                    controller_number=controllers[sat].mdb_index, controller_channel=channel)
                cls.stations[code] = station
                # Compact UTC midnight is inside the operator's local window,
                # but outside the old code's unconverted local-string window.
                MaxicomRuntime.objects.create(site=ccu, station=station, station_id_raw=station.pk,
                                              timestamp=start[:8] + '003000', run_time=minutes)
                Zone.objects.create(code=code, name=code, patch=ccu, boundary_points=boundary,
                                    maxicom_runtime=[station.pk])
            first_station = cls.stations[next(iter(values))]
            first_utc, last_utc = runtime_window_bounds(start, end)
            for stamp in (first_utc, last_utc):
                MaxicomRuntime.objects.create(site=ccu, station=first_station,
                                              station_id_raw=first_station.pk, timestamp=stamp, run_time=1)
            # Keep the provided total, while testing both inclusive boundaries.
            middle = MaxicomRuntime.objects.get(station=first_station, timestamp=start[:8] + '003000')
            middle.run_time -= 2
            middle.save(update_fields=['run_time'])
            for stamp, delta in ((first_utc, -1), (last_utc, 1)):
                outside = (datetime.strptime(stamp, '%Y%m%d%H%M%S') + timedelta(seconds=delta))
                MaxicomRuntime.objects.create(site=ccu, station=first_station,
                                              station_id_raw=first_station.pk,
                                              timestamp=outside.strftime('%Y%m%d%H%M%S'), run_time=99991)

    def setUp(self):
        self.client.force_login(self.user)

    def _params(self, case):
        ccu, start, end, _ = case
        return {'ccu': self.ccus[ccu].pk, 'from': start, 'to': end}

    def test_all_three_operator_examples_match_the_web_matrix(self):
        for case in CASES:
            with self.subTest(ccu=case[0]):
                response = self.client.get('/irrigation/', {**self._params(case), 'format': 'json'})
                self.assertEqual(response.status_code, 200)
                data = response.json()
                actual = {}
                for row in data['rows']:
                    for column, value in zip(data['controllers'], row['values']):
                        if value:
                            actual[f'{column.removeprefix("SAT ")}-{row["channel"]}'] = value
                self.assertEqual(actual, case[3])
                self.assertTrue(data['flow_data']['notice'])

    def test_excel_matches_the_same_window_and_cell_values(self):
        from core.excel_exports import build_irrigation_report_excel
        for ccu, start, end, expected in CASES:
            _, buffer = build_irrigation_report_excel(start, end, [self.ccus[ccu].pk])
            workbook = load_workbook(io.BytesIO(buffer.getvalue()), data_only=True)
            sheet = workbook[f'CCU{ccu}_{self.ccus[ccu].name}'[:31]]
            actual = {}
            for column in range(2, sheet.max_column + 1):
                satellite = str(sheet.cell(4, column).value).removeprefix('SAT ')
                for channel in range(1, 25):
                    value = sheet.cell(4 + channel, column).value
                    if value:
                        actual[f'{satellite}-{channel}'] = value
            self.assertEqual(actual, expected)

    def test_heatmap_uses_the_same_local_window(self):
        response = self.client.get('/api/irrigation/zone-heatmap/', self._params(CASES[0]))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        actual = {z['code']: z['runtime_minutes'] for z in data['zones'] if z['runtime_minutes']}
        self.assertEqual(actual, CASES[0][3])

    def test_complete_day_flow_reads_only_its_local_day(self):
        ccu = self.ccus[7]
        flow = MaxicomFlowZone.objects.create(site=ccu, mdb_index=900001, name='Meter', join_site=True)
        for day, liters in [('20261023', 228507), ('20261024', 7420)]:
            MaxicomFlowDaily.objects.create(zone=flow, site=ccu, flow_day=day, liters=liters)
        response = self.client.get('/irrigation/', {
            'ccu': ccu.pk, 'from': '202610232200', 'to': '202610242159', 'format': 'json'})
        data = response.json()
        self.assertEqual(data['flow_data']['sites'][0]['value'], 7420)
        self.assertEqual(data['flow_data']['notice'], '')

    def test_invalid_local_window_returns_bad_request(self):
        self.assertEqual(self.client.get('/irrigation/', {'from': 'bad'}).status_code, 400)

    def test_second_precision_survives_initial_page_render(self):
        response = self.client.get('/irrigation/', self._params(CASES[2]))
        self.assertContains(response, 'value="2026-08-18T07:01:00" step="1"')
        self.assertContains(response, 'value="2026-08-19T06:57:59" step="1"')

    def test_latest_station_uses_its_local_irrigation_day(self):
        from core.views import _format_irrig_ts, _latest_runtime_by_station
        station = self.stations['7-4-24']
        MaxicomRuntime.objects.filter(station=station).delete()
        for ts, minutes in [('20261023135959', 99), ('20261023140000', 10),
                            ('20261023233000', 20), ('20261024063000', 30)]:
            MaxicomRuntime.objects.create(site=self.ccus[7], station=station,
                                          station_id_raw=station.pk, timestamp=ts, run_time=minutes)
        value = _latest_runtime_by_station([station.pk])[station.pk]
        self.assertEqual(value['minutes'], 60)
        self.assertEqual(_format_irrig_ts(value['ts']), '2026-10-24 14:30')


@override_settings(TIME_ZONE='Asia/Shanghai')
class MaxicomImportTests(TestCase):
    module = 'core.management.commands.import_maxicom_mdb_linux'

    def _run_import(self, flow_exit=0):
        tables = {
            'SITE_CF': [{'IndexNumber': '7', 'SiteNumber': '7', 'IndexName': 'Site Seven'}],
            'CTROL_CF': [{'IndexNumber': '900707', 'IndexName': 'SAT 7-1',
                          'ControllerSiteNumber': '7', 'ControllerLinkChannel': '1'}],
            'STATN_CF': [{'IndexNumber': '900701', 'StationSiteNumber': '7',
                          'StationControllerNumber': '900707', 'StationControllerChannel': '1',
                          'DateOpen': '20260801000000', 'DateClose': '', 'StationFlowFactor': '2'}],
            'XA_LOG': [{'XactStamp': '20260901225400', 'XactIndex': '900701',
                        'SignalTable': 'S', 'SignalType': 'M', 'SignalValue': '20'}],
            'FLOZO_CF': [{'IndexNumber': '900777', 'FlowZoneSiteNumber': '7', 'IndexName': 'Meter'}],
        }
        process = Mock()
        process.stdout = io.StringIO('XactStamp|XactIndex|FlowZoneValue|FlowZoneMultiplier|SiteID\n'
                                    '20260901135959|900777|3|2|7\n'
                                    '20260901140000|900777|4|2|7\n')
        process.wait.return_value = flow_exit
        with patch(self.module + '._mdb_rows', side_effect=lambda _path, table: iter(tables[table])) as exported, \
                patch(self.module + '.os.path.exists', return_value=True), \
                patch(self.module + '.subprocess.Popen', return_value=process), \
                patch('core.sensor_flow.import_sensor_flow'):
            call_command('import_maxicom_mdb_linux', mdb='fixture.mdb', stdout=io.StringIO())
        self.assertNotIn('XA_RuntimeProject', [call.args[1] for call in exported.call_args_list])

    def test_import_uses_actual_utc_logs_and_local_flow_days(self):
        Patch.objects.get_or_create(code='CCU7', defaults={'name': 'Area 7'})
        self._run_import()
        self.assertEqual(list(MaxicomRuntime.objects.values_list('timestamp', 'run_time')),
                         [('20260901225400', 20)])
        self.assertEqual(dict(MaxicomFlowDaily.objects.values_list('flow_day', 'liters')),
                         {'20260901': 6, '20260902': 8})

    def test_failed_meter_export_preserves_existing_daily_records(self):
        ccu, _ = Patch.objects.get_or_create(code='CCU7', defaults={'name': 'Area 7'})
        meter = MaxicomFlowZone.objects.create(site=ccu, mdb_index=123456, name='Previous')
        previous = MaxicomFlowDaily.objects.create(site=ccu, zone=meter, flow_day='20260901', liters=123)
        with self.assertRaises(CommandError):
            self._run_import(flow_exit=1)
        self.assertEqual(MaxicomFlowDaily.objects.get(pk=previous.pk).liters, 123)


@override_settings(TIME_ZONE='Asia/Shanghai', CACHES={
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}})
class CrossNamedControllerTests(TestCase):
    """LR37-4 belongs to source key 12 / displayed Site 14, not CCU37."""
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser('lr-runtime-review', '', 'test-only')
        cls.ccu14, _ = Patch.objects.get_or_create(code='CCU14', defaults={'name': 'Parking'})
        cls.ccu37, _ = Patch.objects.get_or_create(code='CCU37', defaults={'name': 'West Loop'})
        cls.old_station = Patch.objects.create(
            code='station-4421', mdb_index=4421, parent=cls.ccu37,
            name='LR valve', site_number=12, controller_number=213, controller_channel=1)
        cls.zone = Zone.objects.create(code='14-7-1', name='Parking LR zone', patch=cls.ccu14)
        cls.run_import()
        call_command('populate_zone_maxicom_runtime', stdout=io.StringIO())

    @classmethod
    def run_import(cls):
        tables = {
            'SITE_CF': [{'IndexNumber': '12', 'SiteNumber': '14'},
                        {'IndexNumber': '15', 'SiteNumber': '37'}],
            'CTROL_CF': [{'IndexNumber': '213', 'IndexName': 'LR37-4',
                          'ControllerSiteNumber': '12', 'ControllerLinkNumber': '25',
                          'ControllerLinkChannel': '7'},
                         {'IndexNumber': '182', 'IndexName': 'SAT7(CCHP)',
                          'ControllerSiteNumber': '15', 'ControllerLinkNumber': '28',
                          'ControllerLinkChannel': '7'}],
            'STATN_CF': [{'IndexNumber': str(4420 + ch), 'StationSiteNumber': '12',
                          'StationControllerNumber': '213', 'StationControllerChannel': str(ch)}
                         for ch in range(1, 25)] + [
                            {'IndexNumber': '8001', 'StationSiteNumber': '15',
                             'StationControllerNumber': '182', 'StationControllerChannel': '1'}],
            'XA_LOG': [{'XactStamp': '20261017222147', 'XactIndex': '4421',
                        'SignalTable': 'S', 'SignalType': 'M', 'SignalValue': '29'},
                       {'XactStamp': '20261017223000', 'XactIndex': '4421',
                        'SignalTable': 'S', 'SignalType': 'M', 'SignalValue': '58'},
                       {'XactStamp': '20261017223000', 'XactIndex': '8001',
                        'SignalTable': 'S', 'SignalType': 'M', 'SignalValue': '15'}],
        }
        module = 'core.management.commands.import_maxicom_mdb_linux'
        with patch(module + '._mdb_rows', side_effect=lambda _path, table: iter(tables[table])), \
                patch(module + '.os.path.exists', return_value=True), patch(module + '.Command._import_flow'):
            call_command('import_maxicom_mdb_linux', mdb='fixture.mdb', stdout=io.StringIO())

    def setUp(self):
        self.client.force_login(self.user)
        self.params = {'from': '20261018000000', 'to': '20261018235959'}

    def test_source_membership_repairs_existing_ids_and_remains_stable_on_reimport(self):
        for _ in range(2):
            station = Patch.objects.get(code='station-4421')
            self.assertEqual(station.pk, self.old_station.pk)
            self.assertEqual(station.parent_id, self.ccu14.pk)
            self.assertEqual(Patch.objects.filter(controller_number=213, parent=self.ccu14).count(), 24)
            controller = MaxicomController.objects.get(mdb_index=213)
            self.assertEqual((controller.name, controller.site_id, controller.link_channel),
                             ('LR37-4', self.ccu14.pk, 7))
            self.assertEqual(list(MaxicomRuntime.objects.filter(station=station)
                                  .values_list('site_id', 'run_time')),
                             [(self.ccu14.pk, 29), (self.ccu14.pk, 58)])
            self.run_import()

    def test_all_and_single_ccu_web_tables_count_lr_exactly_once(self):
        all_data = self.client.get('/irrigation/', {**self.params, 'format': 'json'}).json()
        tables = {t['ccu_code']: t for t in all_data['ccu_tables']}
        for number in (14, 37):
            single = self.client.get('/irrigation/', {**self.params, 'format': 'json',
                                                     'ccu': getattr(self, f'ccu{number}').pk}).json()
            self.assertEqual(single['controllers'], tables[f'CCU{number}']['controllers'])
            self.assertEqual(single['grand_total'], 87 if number == 14 else 15)
        self.assertIn('LR37-4', tables['CCU14']['controllers'])
        self.assertNotIn('LR37-4', tables['CCU37']['controllers'])
        self.assertEqual(sum(t['grand_total'] for t in tables.values()), 102)
        html = self.client.get('/irrigation/', {**self.params, 'ccu': self.ccu14.pk})
        self.assertContains(html, '<th>LR37-4</th>', html=True)

    def test_excel_uses_the_same_column_and_minutes(self):
        from core.excel_exports import build_irrigation_report_excel
        _, data = build_irrigation_report_excel(self.params['from'], self.params['to'], [self.ccu14.pk])
        workbook = load_workbook(io.BytesIO(data.getvalue()), data_only=True)
        sheet = workbook[f'CCU14_{self.ccu14.name}'[:31]]
        column = next(cell.column for cell in sheet[4] if cell.value == 'LR37-4')
        self.assertEqual(sheet.cell(5, column).value, 87)

    def test_zone_mapping_and_filtered_heatmap_follow_real_link_seven(self):
        self.zone.refresh_from_db()
        self.assertEqual(self.zone.maxicom_runtime, [self.old_station.pk])
        data = self.client.get('/api/irrigation/zone-heatmap/',
                               {**self.params, 'ccu': self.ccu14.pk}).json()
        self.assertEqual([(z['code'], z['runtime_minutes']) for z in data['unmapped']], [('14-7-1', 87)])
        other = self.client.get('/api/irrigation/zone-heatmap/',
                                {**self.params, 'ccu': self.ccu37.pk}).json()
        self.assertFalse(other['zones'] or other['unmapped'])


@override_settings(TIME_ZONE='Asia/Shanghai', CACHES={
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}})
class ExtendedControllerChannelTests(TestCase):
    """CCU6 has a real 40-channel satellite; channels 25–40 must survive reports."""
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser('extended-runtime-review', '', 'test-only')
        cls.ccu, _ = Patch.objects.get_or_create(code='CCU6', defaults={'name': 'CCU Six'})
        cls.ctrl = MaxicomController.objects.create(
            site=cls.ccu, mdb_index=601, name='SAT6-7 (40stn, 6-8)',
            site_number=6, link_number=1, link_channel=7)
        for channel in range(1, 41):
            station = Patch.objects.create(
                code=f'extended-station-{channel}', parent=cls.ccu,
                controller_number=cls.ctrl.mdb_index, controller_channel=channel)
            if channel in (1, 24, 25, 40):
                MaxicomRuntime.objects.create(
                    site=cls.ccu, station=station, station_id_raw=station.pk,
                    timestamp='20261017223000', run_time=channel)

    def setUp(self):
        self.client.force_login(self.user)
        self.params = {'from': '20261018000000', 'to': '20261018235959', 'ccu': self.ccu.pk}

    def test_extended_channels_and_totals_match_in_all_and_single_ccu_tables(self):
        single = self.client.get('/irrigation/', {**self.params, 'format': 'json'}).json()
        all_params = {k: v for k, v in self.params.items() if k != 'ccu'}
        all_data = self.client.get('/irrigation/', {**all_params, 'format': 'json'}).json()
        table = next(t for t in all_data['ccu_tables'] if t['ccu_id'] == self.ccu.pk)
        for data in (single, table):
            self.assertEqual(data['grand_total'], 90)
            self.assertEqual([r['channel'] for r in data['rows']], list(range(1, 41)))
            self.assertEqual({r['channel']: r['values'][0] for r in data['rows'] if r['total']},
                             {1: 1, 24: 24, 25: 25, 40: 40})
        html = self.client.get('/irrigation/', self.params)
        self.assertContains(html, '>40</td>', html=False)

    def test_configured_extended_channels_still_appear_in_an_empty_time_window(self):
        data = self.client.get('/irrigation/', {**self.params, 'format': 'json',
                                               'from': '20261019000000', 'to': '20261019235959'}).json()
        self.assertEqual(len(data['rows']), 40)
        self.assertEqual(data['grand_total'], 0)

    def test_excel_retains_channel_40_and_its_minutes(self):
        from core.excel_exports import build_irrigation_report_excel
        _, data = build_irrigation_report_excel(self.params['from'], self.params['to'], [self.ccu.pk])
        workbook = load_workbook(io.BytesIO(data.getvalue()), data_only=True)
        sheet = workbook[f'{self.ccu.code}_{self.ccu.name}'[:31]]
        self.assertEqual(sheet['A44'].value, '40')
        self.assertEqual(sheet['B44'].value, 40)
        self.assertEqual(sum(sheet.cell(row, 2).value or 0 for row in range(5, 45)), 90)

    def test_pdf_can_render_the_extended_matrix(self):
        response = self.client.get('/irrigation/report.pdf', self.params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF-'))
