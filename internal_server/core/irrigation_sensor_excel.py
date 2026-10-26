"""Physical-sensor sheet for the existing irrigation Excel/email report."""
from collections import defaultdict
from datetime import datetime

from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from core.import_status import status as import_status
from core.irrig_time import _compact_datetime
from core.models import MaxicomController
from core.sensor_flow import sensor_window


def add_sensor_sheet(workbook, date_from, date_to, ccu_ids=None):
    """Keep source Sites independent, including whole meters shared by CCUs."""
    data = sensor_window(date_from, date_to)
    associations = defaultdict(set)
    selected_sources = set()
    wanted = {int(v) for v in (ccu_ids or []) if str(v).strip().isdigit()}
    for source, ccu_id, code in MaxicomController.objects.values_list('site_number', 'site_id', 'site__code'):
        associations[source].add(code)
        if ccu_id in wanted:
            selected_sources.add(source)
    sites = [s for s in data['sites'] if not ccu_ids or int(s['key']) in selected_sources]
    sync = import_status()
    ws = workbook.create_sheet('传感器水量', 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '1B4332'

    ws.merge_cells('A1:G1')
    ws['A1'] = '实际传感器水量'
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color='1B4332')
    ws.row_dimensions[1].height = 28
    ws['A2'] = '查询开始'
    ws.merge_cells('B2:C2')
    ws['B2'] = _compact_datetime(date_from) if date_from else '未指定'
    ws['D2'] = '查询结束'
    ws.merge_cells('E2:G2')
    ws['E2'] = _compact_datetime(date_to, end=True) if date_to else '未指定'
    ws['A3'] = '最近同步'
    ws.merge_cells('B3:C3')
    ws['B3'] = (timezone.localtime(datetime.fromisoformat(sync['last_heartbeat'])).replace(tzinfo=None)
                if sync.get('last_heartbeat') else '暂无同步')
    ws['D3'] = '数据备份'
    ws.merge_cells('E3:G3')
    ws['E3'] = sync.get('source') or '暂无记录'
    for cell in ('B2', 'E2', 'B3'):
        ws[cell].number_format = 'yyyy-mm-dd hh:mm:ss'
    ws.merge_cells('A4:G4')
    ws['A4'] = '上海时间（24小时制）。来源：Rain Bird 实际流量传感器。总表、分表分别列示，不重复相加。'
    ws.merge_cells('A5:G5')
    warning = data.get('notice') or data.get('coverage_note')
    if not sites and not warning:
        warning = '当前 CCU 范围没有已关联的传感器站点。'
    ws['A5'] = warning or '0 表示备份覆盖时段内无流量日志；— 表示无法确认。共用站点保留整表读数，不按 CCU 分摊。'
    ws['A5'].fill = PatternFill('solid', fgColor='FFF2D5' if warning else 'F2F6F3')

    headers = ['Site', '源站点名称', '通道', '传感器名称', '用水量（升）', '记录状态', '关联 CCU']
    for col, label in enumerate(headers, 1):
        cell = ws.cell(7, col, label)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1B4332')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(right=Side(style='thin', color='FFFFFF'))
    ws.row_dimensions[7].height = 26

    row = 8
    for group, site in enumerate(sites):
        codes = sorted(associations.get(int(site['key']), ()), key=lambda c: (len(c), c))
        for offset, sensor in enumerate(site['sensors']):
            state = ('部分记录' if sensor['value'] is not None else '未覆盖') if not sensor['complete'] else (
                    '已停用' if not sensor['enabled'] else '实际记录')
            values = [site['number'], site['name'], sensor['channel'], sensor['name'],
                      sensor['value'] if sensor['value'] is not None else '—', state,
                      '、'.join(codes) or '未关联']
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, value)
                if isinstance(value, str):
                    cell.data_type = 's'  # Source labels must remain literal text.
                cell.font = Font(name='Arial', size=10, color='243A30')
                cell.fill = PatternFill('solid', fgColor='F0F5F1' if group % 2 == 0 else 'FFFFFF')
                cell.alignment = Alignment(horizontal='right' if col == 5 else 'center' if col in (1, 3) else 'left',
                                           vertical='center', wrap_text=col in (2, 4, 7))
                if offset == 0:
                    cell.border = Border(top=Side(style='thin', color='D3DED6'))
            value = sensor['value']
            ws.cell(row, 5).number_format = '#,##0' if value is None or float(value).is_integer() else '#,##0.00'
            if not sensor['complete']:
                ws.cell(row, 6).font = Font(name='Arial', size=10, color='9A5A0A')
            ws.row_dimensions[row].height = 34 if len(sensor['name']) > 30 or len(site['name']) > 29 or len(codes) > 2 else 24
            row += 1

    for line in (2, 3, 4, 5):
        ws.row_dimensions[line].height = 28 if line in (4, 5) else 23
        for cell in ws[line]:
            if cell.value is not None:
                cell.font = Font(name='Arial', size=10, color='56675C')
                cell.alignment = Alignment(vertical='center', wrap_text=line in (4, 5))
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    for col, width in zip('ABCDEFG', (10, 32, 8, 36, 17, 16, 23)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A8'
    ws.auto_filter.ref = f'A7:G{max(7, row - 1)}'
    ws.print_title_rows = '1:7'
    ws.print_area = f'A1:G{max(7, row - 1)}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    ws.oddFooter.center.text = '第 &P 页，共 &N 页'
    workbook.active = 0
    return ws
