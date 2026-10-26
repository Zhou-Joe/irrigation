"""Scheduled XLSX water readings share web windows and preserve source identity."""
import io
from collections import defaultdict
from datetime import date, datetime, timezone as dt_timezone
from unittest.mock import patch

from django.core import mail
from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from core.excel_exports import build_irrigation_report_excel
from core.flow_data import SAMPLE, encode_samples, sample_offset_for_hour
from core.models import (EmailReportConfig, MaxicomController, MaxicomFlowSensor,
                         MaxicomSensorDaily, Patch)
from core.sensor_flow import pulse_microliters
from core.tests_sensor_flow import SIGNALS


@override_settings(TIME_ZONE='Asia/Shanghai', EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class IrrigationSensorExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ccu3,_=Patch.objects.get_or_create(code='CCU3',defaults={'name':'TL'})
        cls.ccu5,_=Patch.objects.get_or_create(code='CCU5',defaults={'name':'FL'})
        for idx,ccu,source in [(980001,cls.ccu3,5),(980002,cls.ccu5,3)]:
            MaxicomController.objects.create(mdb_index=idx,name=f'SAT {ccu.code[3:]}-1',
                site=ccu,site_number=source,link_number=1,link_channel=1)
        cls.sensors={}
        for idx,channel,name in [(82,15,'Low'),(83,16,'High'),(84,17,'WP')]:
            cls.sensors[idx]=MaxicomFlowSensor.objects.create(mdb_index=idx,name=f'{channel}: Flow Sensor {name}',
                source_site_index=5,source_site_number=3,source_site_name='3 Tomorrowland',channel=channel,
                coverage_start='20261023000000',coverage_end='20261025235959')
        totals=defaultdict(lambda:[0,0,bytearray()])
        for r in SIGNALS['sensors']:
            if r['SignalType']!='F':continue
            day,base=sample_offset_for_hour(r['XactStamp'][:10])
            volume=pulse_microliters(r['SignalValue'],r['SignalMultiplier'])
            bucket=totals[(day,int(r['XactIndex']))]
            bucket[0]+=volume;bucket[1]+=1
            bucket[2].extend(SAMPLE.pack(base+int(r['XactStamp'][10:12])*60+int(r['XactStamp'][12:14]),volume))
        for (day,idx),(volume,count,samples) in totals.items():
            MaxicomSensorDaily.objects.create(sensor=cls.sensors[idx],flow_day=day,microliters=volume,
                sample_count=count,samples=encode_samples(samples))

    def workbook(self,start='20261024153200',end='20261025153200',ccus=None):
        _,buffer=build_irrigation_report_excel(start,end,ccu_ids=ccus)
        return load_workbook(io.BytesIO(buffer.getvalue()))

    def readings(self,book):
        return {r[3].value:r[4].value for r in book['传感器水量'].iter_rows(min_row=8)}

    def test_original_operator_sensor_values_are_numeric_and_share_exact_runtime_window(self):
        book=self.workbook(ccus=[self.ccu3.pk]);ws=book['传感器水量']
        self.assertEqual(book.active.title,'传感器水量')
        self.assertEqual(self.readings(book),{'15: Flow Sensor Low':2380,'16: Flow Sensor High':0,'17: Flow Sensor WP':34200})
        self.assertEqual(ws['B2'].value,datetime(2026,10,24,15,32))
        self.assertEqual(ws['E2'].value,datetime(2026,10,25,15,32))
        self.assertEqual(ws['E8'].data_type,'n')
        self.assertEqual(ws.freeze_panes,'A8');self.assertEqual(ws.auto_filter.ref,'A7:G10')
        self.assertEqual(ws['G8'].value,'CCU3')
        self.assertNotIn('Sheet',book.sheetnames)
        self.assertIn(f'CCU3_{self.ccu3.name}'[:31],book.sheetnames)

    def test_filtered_ccu_uses_internal_site_key_and_keeps_shared_meters_whole(self):
        ccu37,_=Patch.objects.get_or_create(code='CCU37',defaults={'name':'Shared'})
        ccu14,_=Patch.objects.get_or_create(code='CCU14',defaults={'name':'Parking'})
        for idx,ccu in [(980003,ccu37),(980004,ccu14)]:
            MaxicomController.objects.create(mdb_index=idx,name=f'SAT {ccu.code[3:]}-2',site=ccu,
                site_number=12,link_number=1,link_channel=2)
        MaxicomFlowSensor.objects.create(mdb_index=98,name='Shared FS',source_site_index=12,
            source_site_number=14,source_site_name='Shared Site',channel=15,
            coverage_start='20261023000000',coverage_end='20261025235959')
        self.assertEqual(self.readings(self.workbook(ccus=[self.ccu5.pk])),{})
        book=self.workbook(ccus=[ccu37.pk]);ws=book['传感器水量']
        self.assertEqual(self.readings(book),{'Shared FS':0})
        self.assertEqual(ws['G8'].value,'CCU14、CCU37')
        self.assertNotIn('CCU3',ws['G8'].value.split('、'))

    def test_missing_coverage_and_disabled_meters_do_not_turn_into_zero_liters(self):
        book=self.workbook('20261027153200','20261028153200')
        self.assertEqual(set(self.readings(book).values()),{'—'})
        self.assertEqual(book['传感器水量']['F8'].value,'未覆盖')
        MaxicomFlowSensor.objects.filter(mdb_index=83).update(is_enabled=False)
        book=self.workbook();self.assertEqual(self.readings(book)['16: Flow Sensor High'],'—')
        self.assertEqual(book['传感器水量']['F9'].value,'已停用')

    def test_sensor_labels_are_literal_excel_text_and_unknown_sync_is_explicit(self):
        MaxicomFlowSensor.objects.filter(mdb_index=82).update(name='=1+1',source_site_name='=HYPERLINK("https://example.invalid")')
        ws=self.workbook()['传感器水量']
        self.assertEqual(ws['D8'].value,'=1+1');self.assertEqual(ws['D8'].data_type,'s')
        self.assertEqual(ws['B8'].data_type,'s');self.assertEqual(ws['B3'].value,'暂无同步')

    def test_daily_dispatch_attaches_sensor_sheet_using_existing_7am_to_65959_window(self):
        from django.core.mail import get_connection
        cfg=EmailReportConfig.objects.create(name='每日推送',frequency='daily',send_hour=9,
            report_types=['irrigation'],recipients='fixture@example.invalid',
            report_windows={'irrigation':{'days_start':1,'days_end':0,'hour_start':7,'hour_end':6}})
        now=datetime(2026,10,26,1,0,tzinfo=dt_timezone.utc)
        with patch('core.report_export.timezone.now',return_value=now), \
                patch('core.smtp_utils.is_smtp_configured',return_value=True), \
                patch('core.smtp_utils.get_email_connection',return_value=get_connection('django.core.mail.backends.locmem.EmailBackend')), \
                patch('core.smtp_utils.resolve_from_email',return_value='sender@example.invalid'):
            call_command('send_report_email',config_id=cfg.pk,stdout=io.StringIO())
        self.assertEqual(len(mail.outbox),1)
        attachment=mail.outbox[0].attachments[0]
        book=load_workbook(io.BytesIO(attachment.content));ws=book['传感器水量']
        self.assertEqual(ws['B2'].value,datetime(2026,10,25,7))
        self.assertEqual(ws['E2'].value,datetime(2026,10,26,6,59,59))
        self.assertEqual(len(self.readings(book)),3)
